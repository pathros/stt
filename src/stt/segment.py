from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import json


@dataclass(slots=True)
class Segment:
    """Un fragmento de transcripción con marcas de tiempo (en segundos)."""

    start: float

    end: float

    text: str

    # Vacío si no se separaron voces. Con --speakers: "Hablante 1", un nombre, etc.
    speaker: str | None = None


def _fmt_timestamp(seconds: float, sep: str) -> str:
    """Formatea segundos como HH:MM:SS<sep>mmm (sep ',' para SRT, '.' para VTT)."""

    if seconds < 0:

        seconds = 0.0

    ms = int(round(seconds * 1000))

    hours, ms = divmod(ms, 3_600_000)

    minutes, ms = divmod(ms, 60_000)

    secs, ms = divmod(ms, 1000)

    return f"{hours:02d}:{minutes:02d}:{secs:02d}{sep}{ms:03d}"


@dataclass(slots=True)
class _Block:
    """Párrafo ya agrupado. speaker es None cuando no se separaron voces."""

    speaker: str | None

    text: str


def _blocks(segments: list[Segment], gap: float = 2.0) -> list[_Block]:
    """Agrupa por pausas y, si hay voces, también cuando cambia quién habla."""

    blocks: list[_Block] = []

    current: list[str] = []

    current_speaker: str | None = None

    prev_end: float | None = None

    started = False

    for s in segments:

        text = s.text.strip()

        if not text:

            continue

        gap_break = prev_end is not None and (s.start - prev_end) > gap

        speaker_break = started and s.speaker != current_speaker

        if current and (gap_break or speaker_break):

            blocks.append(_Block(current_speaker, " ".join(current)))

            current = []

        current.append(text)

        current_speaker = s.speaker

        prev_end = s.end

        started = True

    if current:

        blocks.append(_Block(current_speaker, " ".join(current)))

    return blocks


def _paragraphs(segments: list[Segment], gap: float = 2.0) -> list[str]:
    """Agrupa segmentos en párrafos usando las pausas del audio."""

    return [block.text for block in _blocks(segments, gap)]


def write_txt(segments: list[Segment], path: Path) -> None:
    """Texto plano pensado para pegar/abrir en Word.

    Une segmentos en párrafos (por pausas del audio) y separa párrafos
    con línea en blanco. Así Word no hereda un salto por cada frase.
    """

    chunks = []

    for block in _blocks(segments):

        if block.speaker:

            chunks.append(f"{block.speaker}: {block.text}")

        else:

            chunks.append(block.text)

    path.write_text("\n\n".join(chunks) + "\n", encoding="utf8")


def write_docx(
    segments: list[Segment],
    path: Path,
    title: str,
) -> None:
    """Documento Word (.docx) con título y un párrafo por bloque de habla."""

    from docx import Document

    doc = Document()

    doc.add_heading(title, level=1)

    for block in _blocks(segments):

        if block.speaker:

            paragraph = doc.add_paragraph()

            label = paragraph.add_run(f"{block.speaker}: ")

            label.bold = True

            paragraph.add_run(block.text)

        else:

            doc.add_paragraph(block.text)

    doc.save(path)


def write_md(
    segments: list[Segment],
    path: Path,
    title: str,
    timestamps: bool,
) -> None:
    """Markdown con título; opcionalmente con marcas de tiempo por línea."""

    lines = [f"# {title}", ""]

    if timestamps:

        for s in segments:

            if not s.text.strip():

                continue

            ts = _fmt_timestamp(s.start, ".").split(".")[0]

            if s.speaker:

                lines.append(f"**[{ts}] {s.speaker}:** {s.text.strip()}")

            else:

                lines.append(f"**[{ts}]** {s.text.strip()}")

            lines.append("")

    else:

        for block in _blocks(segments):

            if block.speaker:

                lines.append(f"**{block.speaker}:** {block.text}")

            else:

                lines.append(block.text)

            lines.append("")

    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf8")


def write_srt(segments: list[Segment], path: Path) -> None:

    blocks = []

    for i, s in enumerate(segments, start=1):

        start = _fmt_timestamp(s.start, ",")

        end = _fmt_timestamp(s.end, ",")

        text = s.text.strip()

        if s.speaker:

            text = f"{s.speaker}: {text}"

        blocks.append(f"{i}\n{start} --> {end}\n{text}\n")

    path.write_text("\n".join(blocks), encoding="utf8")


def write_vtt(segments: list[Segment], path: Path) -> None:

    lines = ["WEBVTT", ""]

    for s in segments:

        start = _fmt_timestamp(s.start, ".")

        end = _fmt_timestamp(s.end, ".")

        lines.append(f"{start} --> {end}")

        text = s.text.strip()

        if s.speaker:

            text = f"{s.speaker}: {text}"

        lines.append(text)

        lines.append("")

    path.write_text("\n".join(lines), encoding="utf8")


def write_json(segments: list[Segment], path: Path) -> None:

    data = []

    for s in segments:

        item = {"start": s.start, "end": s.end, "text": s.text}

        if s.speaker:

            item["speaker"] = s.speaker

        data.append(item)

    path.write_text(
        json.dumps(data, indent=2, ensure_ascii=False) + "\n",
        encoding="utf8",
    )


SUPPORTED_FORMATS = ("txt", "md", "docx", "srt", "vtt", "json")


def write_output(
    segments: list[Segment],
    out_dir: Path,
    stem: str,
    fmt: str,
    timestamps: bool = False,
    title: str | None = None,
) -> Path:
    """Escribe la transcripción en el formato indicado.

    Devuelve la ruta del archivo escrito.
    """

    out_dir.mkdir(parents=True, exist_ok=True)

    fmt = fmt.lower().lstrip(".")

    title = title or stem

    path = out_dir / f"{stem}.{fmt}"

    if fmt == "txt":

        write_txt(segments, path)

    elif fmt == "md":

        write_md(segments, path, title, timestamps)

    elif fmt == "docx":

        write_docx(segments, path, title)

    elif fmt == "srt":

        write_srt(segments, path)

    elif fmt == "vtt":

        write_vtt(segments, path)

    elif fmt == "json":

        write_json(segments, path)

    else:

        raise ValueError(f"Formato no soportado: {fmt}")

    return path


def write_many(
    segments: list[Segment],
    out_dir: Path,
    stem: str,
    formats: list[str],
    timestamps: bool = False,
    title: str | None = None,
) -> list[Path]:
    """Escribe la transcripción en varios formatos. Devuelve las rutas."""

    return [
        write_output(segments, out_dir, stem, fmt, timestamps, title)
        for fmt in formats
    ]
